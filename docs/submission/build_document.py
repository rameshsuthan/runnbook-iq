"""Build the submission DOCX with the Codex bundled python-docx runtime.

Run from the repository root. Screenshots are genuine captures, never generated UI.
Render the output with the documents skill renderer and inspect every page.
"""
from pathlib import Path
import ast
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/submission/RunbookIQ-Project-Submission.docx'
D = Document()
SEC = D.sections[0]
SEC.page_width, SEC.page_height = Inches(8.5), Inches(11)
SEC.top_margin = SEC.bottom_margin = Inches(.65)
SEC.left_margin = SEC.right_margin = Inches(.7)
SEC.header_distance = SEC.footer_distance = Inches(.3)
WIDTH = 7.1
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    s = D.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.line_spacing = 1.08
D.styles['Normal'].font.size = Pt(10.5)
D.styles['Title'].font.size = Pt(32)
D.styles['Subtitle'].font.size = Pt(14)
D.styles['Subtitle'].font.italic = False
for name, size in [('Heading 1', 19), ('Heading 2', 12.5), ('Heading 3', 11)]:
    D.styles[name].font.size = Pt(size)
    D.styles[name].font.bold = True
    D.styles[name].paragraph_format.space_before = Pt(11)
    D.styles[name].paragraph_format.keep_with_next = True
D.styles['Caption'].font.size = Pt(9)
D.styles['Caption'].font.italic = False
D.styles['Caption'].font.bold = False
D.styles['Caption'].paragraph_format.space_after = Pt(10)
for sty in D.styles:
    if sty.type == 1:
        sty.paragraph_format.widow_control = True
    for border in list(sty.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
footer = SEC.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('RunbookIQ  |  ').font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)
D.core_properties.title = 'RunbookIQ Project Submission'
D.core_properties.subject = 'Architecture application walkthrough and evaluation'
D.core_properties.author = 'rameshsuthan'
D.core_properties.keywords = 'RAG LangGraph Pinecone Gemini Streamlit'
D.core_properties.comments = ''

def p(text='', style=None):
    return D.add_paragraph(text, style)

def h(text, level=1):
    return D.add_heading(text, level)

def page(title):
    heading=h(title); heading.paragraph_format.page_break_before=True

def fig(name, caption, width=WIDTH):
    path = ROOT / ('docs/submission/screenshots/' + name if name[0].isdigit() else 'docs/diagrams/' + name)
    if not path.exists():
        raise FileNotFoundError(path)
    para = p(); para.paragraph_format.keep_with_next = True
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape = para.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set('descr', caption)
    if name == '01-question.png':
        # Native Word crop removes empty browser margins without altering the capture.
        rect=OxmlElement('a:srcRect')
        for edge,value in [('l','16364'),('r','16364'),('t','12222'),('b','35556')]:rect.set(edge,value)
        fill=shape._inline.graphic.graphicData.pic.blipFill
        fill.insert(1,rect)
        shape.height=Inches(width*470/740)
    p(caption, 'Caption')

def table(headers, rows, widths, size=9.4):
    t = D.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths): col.width = Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)
    for ri, values in enumerate([headers, *rows]):
        row = t.rows[0] if ri == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        for ci, (cell, text, width) in enumerate(zip(row.cells, values, widths)):
            cell.width=Inches(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'233C50' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
            para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.03
            run=para.add_run(str(text));run.font.size=Pt(size)
            if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
            if ci==0 and widths[0] < .7:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p().paragraph_format.space_after=Pt(0)
    return t

def code(text):
    para=p();para.paragraph_format.space_after=Pt(9)
    for i,line in enumerate(text.splitlines()):
        run=para.add_run(('\n' if i else '')+line);run.font.name='Courier New';run.font.size=Pt(9)
    return para

def link(label, url):
    para=p(); rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),rel)
    r=OxmlElement('w:r');rpr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'185A78');rpr.append(color);r.append(rpr)
    text=OxmlElement('w:t');text.text=label;r.append(text);hyp.append(r);para._p.append(hyp)

def literal(path, name):
    tree=ast.parse((ROOT/path).read_text())
    for n in tree.body:
        if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id==name for t in n.targets):return ast.literal_eval(n.value)
    raise ValueError(name)

p('THE GEN ACADEMY  ·  WEEK 2 CODE TRACK', 'Subtitle')
p('RunbookIQ', 'Title')
p('Project submission', 'Subtitle')
p('Incident evidence retrieval with LangGraph Pinecone and Gemini')
p('Prepared by rameshsuthan  ·  19 October 2026')
h('What I built')
p('RunbookIQ helps engineers answer documented troubleshooting questions about a fictional order-processing service. It searches nine original runbooks, incident reports and reference guides, then returns cited evidence, recommended checks and specific missing information in a local Streamlit application.')
p('The implemented workflow covers a dense retrieval baseline and optional hybrid search, reranking, multi-query expansion and one bounded output repair. Phase 1 and the revised Phase 2 engineering scope are complete. Nebius is deferred to an optional Phase 3.')
h('What the results establish', 2)
p('All four retrieval configurations found every expected section for the twelve answerable evaluation cases. The final resumed answer experiment matched 12 of 15 expected statuses, but one response denied facts present in its evidence. Retrieval success and valid citation IDs therefore do not establish factual correctness.')
table(['Project dimension','Implemented scope'],[
 ('Audience and surface','Engineers learning incident triage; local Streamlit and CLI'),
 ('Knowledge base','9 synthetic Markdown documents; 21 retained canonical chunks'),
 ('Core stack','Python 3.12, LangGraph, Pinecone, Gemini, Streamlit and SQLite'),
 ('Verification','135 tests at the last engineering gate; 15 recorded evaluation cases'),
 ('Boundary','Documented checks only; no live-system diagnosis or remediation')],[1.7,5.4])
link('Code and project assets on the phase 2 branch','https://github.com/rameshsuthan/runnbook-iq/tree/phase-2')
p('Reading guide: Architecture 2 · Ingestion 3 · Retrieval 4 · LangGraph 5 · Decisions 7 · Prompts 8 · Application screens 9 to 12 · Usage 13 · Evaluation 14 to 16 · Development 17 · Running the project 18', 'Caption')

page('System architecture')
p('The application runs locally. LangGraph coordinates retrieval and generation; Pinecone supplies llama-text-embed-v2 integrated embeddings, dense search and a dedicated reranker. Gemini gemini-3.5-flash-lite produces query alternatives and structured answers. Local canonical records keep evidence consistent across stages.')
fig('system-overview.png','Figure 1. Implemented system boundaries and dependencies. SQLite caches and ledgers are application storage, not conversational memory.')
h('Responsibilities and boundaries',2)
p('Streamlit, the CLI and evaluation runner use the same configuration and composition layer. That layer injects service clients into graph node closures. Request state contains questions, evidence and results; it does not contain credentials or SDK clients.')
p('BM25 keyword search, rank fusion and schema/citation checks run locally. Pinecone receives text for integrated embeddings and canonical metadata; the application does not calculate embedding vectors locally. Gemini receives the question and selected passages. Tracing is disabled.')
p('The result exposes the answer, source passages, effective queries, stage scores, repair count, service operations and usage. There is no deployed API server, background indexer, live observability connection or graph database.')
h('Two phases with a stable evidence contract',2)
p('Phase 1 established canonical chunks, explicit ingestion, dense retrieval, cited generation and evaluation. Phase 2 extended the same candidate and result contracts with lexical retrieval, rank fusion, query variants and repair. Switching the chat route does not require re-embedding; changing the embedding model does.')

page('Corpus and ingestion')
p('The source of truth is data/docs. The documents deliberately contain similar symptoms, different timeout meanings, historical versus current settings and unavailable facts. They contain no production logs or credentials.')
table(['Document group','Files and coverage'],[
 ('Service overview','svc-order-api.md — request path and service boundaries'),
 ('Runbooks','rb-db-pool.md and rb-dependency-timeouts.md — checks and retry limits'),
 ('Historical incidents','incident-db-pool.md and incident-dependency-timeouts.md — findings and timelines'),
 ('References','ref-config.md and ref-errors.md — current limits and error meanings'),
 ('Guides','guide-telemetry.md and guide-triage.md — correlation fields and ambiguity')],[1.5,5.6],9.2)
fig('ingestion-flow.png','Figure 2. Explicit build and ingestion publish a verified active snapshot. Questions never trigger ingestion.')
p('Section-aware chunks target 400 tokens and enforce a 2048-token embedding limit with truncation disabled. IDs, document metadata, revision, content hash and corpus version preserve citation identity. The retained evaluation snapshot contains 21 sections-v1 chunks; the reviewed source splitter now uses sections-v2. A future explicit rebuild changes corpus identity and requires ingestion.',None)
p('A corpus lock, destination binding and pending journal protect synchronization. Batches upsert at most 96 records, remove only obsolete owned IDs, and verify remote freshness before active.json is published atomically. An interrupted journal blocks queries until an explicit resume completes. No scheduled refresh is implemented.')

page('Retrieval and evidence selection')
p('The default baseline searches with the original question. Advanced profiles can retain that question and add at most two validated alternatives. Each query receives up to five dense hits and, when hybrid search is enabled, up to five local BM25 hits.')
fig('retrieval-flow.png','Figure 3. Dense and lexical results share canonical chunks. Fusion produces at most 15 candidates; selection supplies at most 5 passages.')
h('Why combine ranks',2)
p('Dense search captures semantic similarity. BM25 preserves exact lexical signals such as DB_POOL_ACQUIRE_TIMEOUT and configuration names. Reciprocal Rank Fusion adds 1 / (60 + rank) for each unique chunk in each ranked list. It avoids treating embedding and keyword scores as directly comparable. A single dense list retains its original ordering.')
p('The optional bge-reranker-v2-m3 reranker compares the original question with canonical candidate text in one Pinecone inference request. It reorders the existing pool; it does not discover new documents or certify the answer. Returned indices, cardinality and finite scores are checked before mapping selections back to local records.')
h('Consistency and scope',2)
p('Remote IDs, text, hash and version must match the active canonical snapshot. The pipeline rechecks the snapshot before returning the batch. Only user-supplied service and environment filters constrain search; an inferred diagnosis never becomes a filter. Stage scores and query/method provenance remain inspectable.')
p('In the recorded comparison, added retrieval stages produced no measured recall improvement over dense search. The dense profile remains the default; advanced profiles provide controlled experiments and richer inspection.')

page('LangGraph control flow')
p('LangGraph represents the request as explicit stages with conditional exits. The workflow is bounded and sequential. It has eight named nodes; optional features change the work within those nodes and the branches followed.')
fig('langgraph-flow.png','Figure 4. Simplified runtime flow. Errors can stop the graph; successful fresh responses pass validation, with at most one repair attempt.')
h('State nodes and edges',2)
p('GraphState is a typed request record. It carries the question, explicit filters, query variants, candidate batch, selected evidence, provider response, result, operation records, timings and repair counter. A node is a Python function that reads this state and returns a partial update. An edge decides which function runs next.')
p('StateGraph.compile builds the runnable workflow. Each invoke begins a fresh question. No checkpointer or cross-request conversational memory is configured; Streamlit session state only retains the last displayed result. SQLite caches persist separately from the graph.')
h('The repair block',2)
p('A categorized schema or citation error in a fresh answer can follow validate → repair → validate when repair is enabled and the counter is zero. The repair receives the same question and evidence, the rejected response as untrusted data, and a validation category. It never repeats retrieval or expansion. The counter is incremented before the call; another invalid response ends explicitly.')
p('Transport errors and invalid cached answers do not trigger repair. The USD3 guard also applies to repair. This loop checks output structure and citation identity, not semantic truth; it did not detect the q08 factual error. END means the workflow stopped, not that an answer has been proven correct.')

page('Node contracts and decisions')
table(['Node','Work and state update','External work'],[
 ('prepare_request','Validate the question and explicit filters; reset per-request collections and counters.','None'),
 ('prepare_queries','Use the original question or cache/validate up to two alternatives. Return effective queries and warnings.','0 or 1 Gemini call'),
 ('retrieve','Search each query, verify canonical hits, run optional BM25 and fuse lists. Accumulate operations.','Up to 3 dense searches'),
 ('select_evidence','Select at most five chunks; validate optional reranker output.','0 or 1 rerank call'),
 ('fallback','Return insufficient_evidence when answer mode has no selected passages.','None'),
 ('generate','Check the evidence-bound answer cache; otherwise request one structured answer.','0 or 1 Gemini call'),
 ('validate','Check schema, answer sections and citation identity; cache valid results.','None'),
 ('repair','Attempt one correction using unchanged evidence and the validation error.','0 or 1 Gemini call')],[1.3,4.35,1.45],9.2)
h('Conditional exits',2)
p('Invalid input ends as invalid_request. Retrieval or selection failure ends as retrieval_error. Retrieval-only mode ends as retrieved, including when evidence is empty. In answer mode, no evidence takes the deterministic fallback; selected evidence proceeds to cache lookup or generation.')
p('A valid cached answer is revalidated and returned with current request operations. An invalid cached answer ends as invalid_output without regeneration. Provider/setup failures and budget denial remain explicit. Fresh valid output ends as answered or insufficient_evidence; only eligible fresh validation failures can enter repair.')
h('Failure behavior is part of the design',2)
p('A fresh query-expansion failure normally retains the original question with a warning and attempted usage. Required-cache mode instead stops before search when its cache requirement cannot be met. Lexical and token guards reduce unsafe rewrites, but they do not prove semantic equivalence.')
p('SDK retries are disabled. At most one expansion, one answer and one repair can occur per question. The repair counter provides the loop bound; the graph invocation also has a recursion limit of 16 as a backup.')

page('Design decisions and configuration')
table(['Decision','Reason and tradeoff'],[
 ('Synthetic incident corpus','Allows realistic ambiguity and historical/current distinctions without disclosing internal operational information. Limits the strength of generalization claims.'),
 ('Explicit LangGraph workflow','Makes node transitions, abstention and bounded repair inspectable. Adds orchestration code but supports learning and failure analysis.'),
 ('Integrated Pinecone embeddings','Keeps document and query embedding paths aligned. Requires hosted service access and independent quota accounting.'),
 ('One canonical snapshot','BM25, dense hits, reranking and citations use the same IDs and text. Inconsistent remote/local records fail visibly.'),
 ('Dense default with optional hybrid profiles','Allows a controlled comparison. No measured recall benefit justifies making the advanced stages mandatory.'),
 ('Dedicated reranker and bounded expansion','Avoids per-chunk chat grading; limits queries and candidates. Relevance scoring and lexical guards are not truth tests.'),
 ('Deterministic validation and one repair','Checks an explicit output contract with bounded usage. Semantic correctness still needs source review.'),
 ('Evidence-bound caches and persistent ledgers','Reuse depends on corpus, route, prompts and evidence identity. Retrieval still runs before answer cache lookup.')],[2.05,5.05],9.1)
h('Profiles are explicit switches',2)
table(['Configuration','Hybrid','Rerank','Multi query','Repair'],[
 ('config.yaml','Off','Off','Off','Off'),
 ('config.phase2-hybrid.yaml','On','Off','Off','Off'),
 ('config.phase2-rerank.yaml','On','On','Off','Off'),
 ('config.phase2-multiquery.yaml','On','On','On','Off'),
 ('config.phase2-repair.yaml','On','On','On','On')],[3.1,1,1,1,1],9.1)
p('The question itself does not turn these switches on. The full Phase 2 profile selects all four; a plain app launch loads the dense default. Restart with the intended profile and submit again to see its configuration. No reingestion is required to switch retrieval profiles.')

page('Prompts and answer contract')
p('The final comparison fixes answer prompt incident-evidence-v3 and schema answer-v3. The following is the implemented answer system prompt from providers.py. The request payload adds the original question and selected passages with exact chunk IDs, document types and revisions.')
prompt=literal('src/runbook_iq/providers.py','SYSTEM_PROMPT')
para=p(prompt);para.paragraph_format.line_spacing=1.02
for r in para.runs:r.font.size=Pt(9.5)
h('Structured response',2)
table(['Field','Meaning'],[
 ('status','answered or insufficient_evidence'),
 ('documented_evidence','Direct response supported by selected passages with inline IDs'),
 ('recommended_checks','Documented next checks with inline IDs; nonempty for an answer'),
 ('missing_information','Specific unavailable facts or observations'),
 ('citations','Unique exact IDs used inline; must match the inline citation set')],[1.8,5.3],9.3)
h('Expansion and repair instructions',2)
p('Expansion asks for at most two alternatives, preserving identifiers, quoted literals, numbers and negation. It forbids added diagnoses or observations and permits no alternatives when a safe rewrite is unclear. Local checks apply after generation and again on cache reads.')
p('Repair asks for a corrected structured answer using only the unchanged selected evidence. It treats the rejected response as untrusted data, forbids invented sources, requires matching inline and listed IDs, and permits an explicit abstention. Exact prompts are retained in providers.py and queries.py.')

# Application pages are assembled from captures of the actual local Streamlit app.
page('Application walkthrough')
p('The form accepts a troubleshooting question and optional explicit filters. Submission runs the graph once. Opening an expander or rerendering the page does not submit another request.')
fig('01-question.png','Figure 5. Actual local Streamlit question form, captured on 19 October 2026.',6.8)
h('Example question',2)
p('What documented checks distinguish database pool exhaustion from an inventory dependency timeout in order-api?')
p('This multi-document question tests whether the assistant separates a 500 ms pool acquisition wait from a 1500 ms inventory request deadline, and identifies the different metrics and traces to inspect. The response describes documented signals; it cannot establish a current incident cause.')

page('Cited answer')
fig('02-answer.png','Figure 6. Answer content from the actual dense baseline request. Inline chunk IDs connect claims to source passages.',6.8)
p('The result separates documented evidence from recommended checks and missing information. A displayed answered status means the output passed structural validation. A reviewer should still open its sources and compare the claims with the passages.')
p('This screenshot request used one dense search and one fresh Gemini generation. The UI reported 1129 input tokens, 541 output tokens, estimated chat cost USD0, one Pinecone read unit and 19 embedding tokens. No repair ran. These are walkthrough observations, separate from the frozen comparison totals.')

page('Source evidence and retrieval inspection')
fig('03-source.png','Figure 7. Expanded canonical source passage showing its exact chunk ID, document type and revision.',6.3)
fig('04-inspection.png','Figure 8. Phase 2 configuration excerpt for the live-root-cause question. Hybrid, reranking and multi-query are enabled.',6.3)
p('Inspection retains candidate scores and query/method provenance, selected evidence and the repair attempt count. Scores indicate relevance, not confidence. Zero repair attempts is normal: repair runs only after an eligible fresh validation failure.')

page('Abstention when evidence is insufficient')
fig('05-abstention.png','Figure 9. Actual result for the unavailable live-root-cause question. The assistant names missing observations instead of inventing a current cause.',6.8)
p('Question: What is the confirmed root cause of our live incident right now? The corpus contains historical examples and documented checks, but no current incident observations. The expected behavior is an explicit insufficient_evidence result with the unavailable information identified.')
p('The screenshot uses the full Phase 2 profile in a separate local session. Expansion and generation were cache hits, with zero new Gemini calls. One dense search used 16 embedding tokens and one read unit; reranking used one unit. Only the original query remained in the cached query set, which is valid even with multi-query enabled. This walkthrough is separate from the fifteen-case evaluation.')

page('Caches usage and timing')
fig('cache-and-usage.png','Figure 10. Answer caching occurs after retrieval and selection. The numerical example is a previously supplied observed request, not a fixed price per question.')
p('Expansion cache identity includes the question, filters, corpus, prompt/schema and route settings. Answer cache identity also includes the retrieval configuration, exact query variants, repair policy and ordered selected chunk IDs and content hashes. Cached answers are revalidated against the current selection.')
h('How to read null and zero',2)
p('A cache hit has attempted=false and no new provider response, so operation-level chat token fields may be null. Current aggregate chat calls and tokens can be zero. Pinecone embedding tokens, read units and rerank units appear under units, not under chat usage. A cached answer can therefore have zero Gemini calls while retrieval still consumes Pinecone allowance.')
p('usage_complete describes chat input/output/cost completeness, not every Pinecone field or optional reasoning-token detail. Unknown attempted usage remains unknown; known_subtotal sums only known values. A zero cache-operation duration is an unmeasured default. The node timer includes lookup work; repeated validation overwrites that node entry, so evaluation request wall time is the appropriate total.')
h('Budget boundaries',2)
p('A persistent atomic ledger guards cumulative estimated paid application usage at USD3 across phases. The configured Gemini route uses the user-confirmed free tier; paid fallback is disabled. Unknown paid outcomes retain their reservation. Separate local counters reserve Pinecone embedding and rerank allowance; read units are monitored from responses. These counters do not establish the account-wide bill or other applications’ consumption.')

page('Evaluation results')
p('The experiment uses fifteen fixed cases spanning exact identifiers, paraphrases, similar symptoms, multiple documents, history, ambiguity and unavailable information. Twelve cases expect an answer; three expect insufficient evidence. All cases have now been observed, including those originally labeled heldout, so these results are regression evidence rather than an unseen test.')
table(['Retrieval configuration','Expected section hit at 5','Expected section recall at 5'],[
 ('Dense only','1.0','1.0'),('Hybrid','1.0','1.0'),('Hybrid and reranking','1.0','1.0'),('Hybrid reranking and multi query','1.0','1.0')],[3.1,2,2],9.5)
p('Hit asks whether at least one expected section was selected; recall asks what fraction of a case’s expected sections was selected. Both are averaged across the twelve answerable cases. Unanswerable cases are excluded. Selected and fused recall were already saturated; no retrieval quality gain is established.', 'Caption')
table(['Final answer experiment','Recorded result'],[
 ('Completed cases','15 across an initial run and one continuation'),
 ('Completed outcomes','9 answered; 6 insufficient_evidence'),
 ('Exact expected status matches','12/15 completed; 11/15 counting first attempts'),
 ('First attempt interruption','q06 retrieval failure; no q06 chat call before continuation'),
 ('Chat usage across both runs','16 calls: 15 generation and 1 repair; 17064 input and 5843 output tokens'),
 ('Estimated chat cost','USD0 on the configured free tier; not an account invoice'),
 ('Completed request latency','Median 2.23 seconds; range 1.62 to 4.45 seconds'),
 ('Last engineering verification','135 tests passed; offline audit preserved all 15 query sets')],[2.3,4.8],9.4)
p('The answer experiment used warm expansion caches and fresh answer generation. Latency excludes client composition and intentional pauses. The q06 failed search has unknown Pinecone units and remains preserved. Live evaluation preceded the final edge-case fixes; the subsequent audit was offline, with no claim of a post-fix live evaluation.')
h('Known answer quality limits',2)
p('q08 incorrectly denies documented correlation fields despite retrieving them. q02 gives the correct deadline but omits the expected caveat that timeout alone does not establish cause. q04 and q11 return useful supported checks with conservative abstention statuses. Exact status agreement is therefore not a faithfulness score. Source judgments are the implementing assistant’s review, not independent human ratings.')

questions=json.loads((ROOT/'eval/questions.json').read_text())
notes={
 'q01':'Supported pool-wait meaning and three metrics.',
 'q02':'Correct 1500 ms deadline; missing no-cause caveat.',
 'q03':'Supported distinction between deadlines and signals.',
 'q04':'Useful clarifications; conservative abstention differs from expected answer.',
 'q05':'Supported retry boundary and outcome/idempotency checks.',
 'q06':'First request failed at retrieval; continuation answered historical findings.',
 'q07':'Supported historical 12 versus current 20 per process.',
 'q08':'Factual failure: denies documented request_id/trace_id correlation fields.',
 'q09':'Appropriate abstention; no secret invented.',
 'q10':'Appropriate abstention; current observations unavailable.',
 'q11':'Supported checks; conservative abstention differs from expected answer.',
 'q12':'Supported 20 connections, 500 ms and 1500 ms.',
 'q13':'Supported route revert 14:12 UTC and recovery 14:20 UTC.',
 'q14':'Supported: missing status does not prove no reservation.',
 'q15':'Appropriate abstention after one recorded repair.'}
abstain={'q04','q08','q09','q10','q11','q15'}
for start,end,title in [(0,8,'Evaluation cases 1 through 8'),(8,15,'Evaluation cases 9 through 15')]:
    page(title)
    p('The questions and expectations remain frozen. “Answer” below means answered; “Abstain” means insufficient_evidence. The actual status is the completed result, with q06’s failed first attempt retained separately.')
    rows=[]
    for q in questions[start:end]:
        ident=q['id'];expected='Answer' if q['expected_status']=='answered' else 'Abstain';actual='Abstain' if ident in abstain else 'Answer'
        rows.append((ident,q['question'],expected+' / '+actual,notes[ident]))
    table(['Case','Question','Expected / actual','Source review observation'],rows,[.48,2.68,1.14,2.8],9.2)
    if start:
        h('Repair evidence and semantic limits',2)
        p('q15 took one real repair call and ended in a structurally valid, source-supported abstention. The final artifact records two chat calls and repair_attempts=1, but does not retain the initial rejected text or error category. No more specific explanation of that initial defect is claimed.')
        p('q08 passed schema and citation checks. Its error demonstrates why retrieving the correct section, selecting a valid ID and supporting the meaning of a claim are separate evaluation questions. No semantic grading agent has been implemented.')

page('Development iterations and learning')
p('I used Codex to translate the approved project brief into small runnable milestones, implement and test the workflow, explain LangGraph concepts, diagnose integration failures and prepare evaluation artifacts. The work used native execution with explicit design gates and an independent branch review.')
table(['Iteration','Change and observed learning'],[
 ('Phase 1 first milestone','Three documents, Pinecone retrieval and one cited LangGraph answer established the end-to-end path.'),
 ('Corpus and evaluation','Expanded to nine documents and fifteen fixed cases. Canonical IDs and explicit failure states made evidence review possible.'),
 ('Answer contract refinement','Citation and telemetry-field failures led to targeted prompt/schema iterations. Failed experiments remained in the record.'),
 ('Phase 2 retrieval comparison','Added BM25/RRF, reranking and bounded query expansion without changing the answer prompt during comparisons. Recall did not improve.'),
 ('Accounting and repair','Added current-request operation records, evidence-bound caching, one repair and inspectable service usage.'),
 ('Independent review and fixes','Fixed whole-token preservation, cached query token-limit checks and lost usage metadata after malformed hits. Six regression cases failed before fixes and passed afterward.')],[1.65,5.45],9.5)
h('What I learned',2)
p('LangGraph makes control flow explicit, but the graph itself does not establish answer quality. A deterministic validator can enforce an output contract without proving entailment. Bounded repair is useful for that contract; a source-supported answer still needs a separate assessment.')
p('A small corpus can saturate retrieval metrics. Adding stages increases service work without guaranteeing better answers. Frozen questions, corpus identities and retained failures make that tradeoff visible. Known usage, unknown usage and cached operations also need distinct representations.')
h('Remaining work',2)
p('Priorities are targeted investigation of q08, completeness work for q02, and evaluation of status behavior for q04/q11 without rewriting frozen expectations. Production authentication, deployment, live-system actions and a semantic judge are outside this project. Nebius is optional Phase 3 and has no live acceptance claim.')
p('The demo narration and diagrams are prepared. Recording, checking viewer access and submitting the final links remain delivery steps. The screenshot walkthrough in this report is documented separately from the frozen evaluation.')

page('Running and reviewing the project')
h('Run the existing configured workspace',2)
p('Use Python 3.12 and the locked uv environment. Keep credentials in the local ignored .env. The configured account tiers are assertions that should be checked for any new account. Do not print keys, clear usage ledgers or rebuild the retained evaluation corpus for a demo.')
code('uv sync --locked\nuv run runbook-iq --config config.phase2-repair.yaml doctor --offline\nuv run pytest -q')
p('Offline doctor does not load .env; a missing shell-key indicator does not prove the local .env lacks credentials. Tests use temporary data and block outbound network calls.')
code('RUNBOOK_IQ_CONFIG=config.phase2-repair.yaml \\\n  uv run streamlit run app.py --server.address 127.0.0.1 \\\n  --browser.gatherUsageStats false')
p('For a new workspace, follow README setup for tokenizer acquisition, corpus build, dry run and explicit ingestion. Those steps can change corpus identity or consume provider allowance; ordinary question submission never ingests. Preserve the existing frozen reports.')
h('Reviewable project assets',2)
base='https://github.com/rameshsuthan/runnbook-iq/blob/phase-2/'
for label,path in [
 ('README setup and application commands','README.md'),
 ('Complete architecture and seven editable Mermaid flows','docs/architecture.md'),
 ('Fifteen frozen questions and expected evidence','eval/questions.json'),
 ('Recorded answer summary and usage counts','eval/results/phase2-final-summary.json'),
 ('Evaluation source review and retained failure analysis','eval/results/phase2-review.md'),
 ('Independent engineering review and regression fixes','docs/phase2-code-review.md'),
 ('Timed demo narration and screen actions','docs/demo/demo-script.md')]:link(label,base+path)
p('The historical review files retain their original chronology, including earlier Nebius and visual-walkthrough pending notes. PROJECT_CONTEXT.md and PROGRESS.md record later scope changes and the current documentation verification. The 135-test result is the last engineering gate, not a newly rerun suite for this report.')
h('Implementation map',2)
p('graph.py owns orchestration; retrieval.py, keyword.py, fusion.py and reranking.py own evidence selection; providers.py and queries.py own model instructions; generation.py validates output; schemas.py defines contracts. cache.py, budget.py and usage.py support reuse and accounting. app.py renders the user interface.')
OUT.parent.mkdir(parents=True,exist_ok=True)
D.save(OUT)
print(OUT)
print('Paragraphs:',len(D.paragraphs),'Tables:',len(D.tables),'Images:',len(D.inline_shapes))
